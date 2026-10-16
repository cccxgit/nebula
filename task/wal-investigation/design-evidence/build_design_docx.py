"""Render the canonical kickoff Markdown as a navigable Chinese Word document."""
from pathlib import Path
import re
import hashlib
import json
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'SPACE-LEVEL-WAL-REQUIREMENTS-DESIGN.md'
OUTPUT = SOURCE.with_suffix('.docx')
FONT = 'Noto Sans CJK SC'

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
section.left_margin, section.right_margin = Cm(2), Cm(2)
section.header_distance, section.footer_distance = Cm(.7), Cm(.7)

def font_settings(style, size, bold=False, latin=None):
    style.font.name = latin or FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), FONT)
    fonts.set(qn('w:ascii'), latin or FONT)
    fonts.set(qn('w:hAnsi'), latin or FONT)
    # Ensure black cannot be overridden by a theme color.
    color = style.element.get_or_add_rPr().find(qn('w:color'))
    if color is not None:
        for key in ('themeColor', 'themeShade', 'themeTint'):
            color.attrib.pop(qn('w:' + key), None)

normal = doc.styles['Normal']
font_settings(normal, 10.5)
normal.paragraph_format.line_spacing = 1.17
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.widow_control = True
for name, size in [('Title', 20), ('Heading 1', 15), ('Heading 2', 12)]:
    s = doc.styles[name]
    font_settings(s, size, True)
    s.paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.keep_together = True
    s.paragraph_format.line_spacing = 1.12
    pp = s.element.find(qn('w:pPr'))
    if pp is not None:
        for border in pp.findall(qn('w:pBdr')):
            pp.remove(border)
for name in ('Header', 'Footer'):
    font_settings(doc.styles[name], 8)
    doc.styles[name].paragraph_format.space_after = Pt(0)
code_style = doc.styles.add_style('Code Block', 1)
font_settings(code_style, 8.2, latin='DejaVu Sans Mono')
code_style.paragraph_format.line_spacing = 1.12
code_style.paragraph_format.space_after = Pt(1)
code_style.paragraph_format.left_indent = Cm(.2)
code_style.paragraph_format.widow_control = True

header = section.header.paragraphs[0]
header.text = 'Nebula Graph 图空间级 WAL 缓存配置需求设计'
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('0.9 评审稿  ·  ')
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)

token = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')

def inline(p, text):
    for part in token.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'DejaVu Sans Mono'
            run.font.size = Pt(9)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
        elif part.startswith('['):
            m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
            if not m:
                p.add_run(part)
                continue
            label, target = m.groups()
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            pr = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '244A67'); pr.append(color)
            fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'), FONT); pr.append(fonts)
            run.append(pr)
            value = OxmlElement('w:t'); value.text = label; run.append(value)
            hyperlink.append(run); p._p.append(hyperlink)
        else:
            p.add_run(part)

def table(rows):
    n = len(rows[0])
    t = doc.add_table(rows=0, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # Identifier/label columns stay narrow, while contract/evidence has room.
    first = rows[0][0]
    if n == 2:
        widths = [5.8, 11.2]
    elif n == 3 and first in ('编号', '工作包', '决策项'):
        widths = [2.4, 6.5, 8.1]
    elif n == 3 and first == '逻辑容量':
        widths = [3.0, 5.5, 8.5]
    else:
        widths = [4.0, 6.3, 6.7]
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + side)
        edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9')
        borders.append(edge)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side, val in [('top', '85'), ('bottom', '85'), ('left', '100'), ('right', '100')]:
        elem = OxmlElement('w:' + side); elem.set(qn('w:w'), val); elem.set(qn('w:type'), 'dxa'); margins.append(elem)
    pr.append(margins)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        rowpr = t.rows[-1]._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); rowpr.append(cant)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); rowpr.append(repeat)
        for ci, (cell, value, width) in enumerate(zip(cells, row, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.12
            if ci == 0 and first in ('编号', '逻辑容量'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9.5)
                if ri == 0:
                    run.bold = True
            if ri == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EEF3'); cell._tc.get_or_add_tcPr().append(shade)
    # A real paragraph supplies separation without a fake blank table row.
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = .3
    p.add_run().font.size = Pt(2)

lines = SOURCE.read_text().splitlines()
i = 0
code = False
keep_code = False
while i < len(lines):
    line = lines[i]
    if line.startswith('```'):
        if not code:
            end = i + 1
            while end < len(lines) and not lines[end].startswith('```'):
                end += 1
            length = end - i - 1
            keep_code = length <= 2 or (line == '```json' and length <= 16)
            if keep_code and doc.paragraphs:
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
            code = True
        else:
            code = False
            if keep_code and doc.paragraphs:
                doc.paragraphs[-1].paragraph_format.keep_with_next = False
                doc.paragraphs[-1].paragraph_format.space_after = Pt(5)
        i += 1
        continue
    if code:
        p = doc.add_paragraph(style='Code Block')
        if keep_code:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
        p.add_run(line if line else ' ')
        i += 1
        continue
    if not line.strip():
        i += 1
        continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            row = [v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[: -]+', v) for v in row):
                rows.append(row)
            i += 1
        table(rows)
        continue
    m = re.match(r'^(#{1,3}) (.+)$', line)
    if m:
        level, title = len(m[1]), m[2]
        p = doc.add_paragraph(style={1:'Title',2:'Heading 1',3:'Heading 2'}[level])
        inline(p, title)
        if level == 2 and title.startswith(('附录 B', '附录 C')):
            p.paragraph_format.page_break_before = True
    else:
        p = doc.add_paragraph()
        inline(p, line)
    i += 1

doc.core_properties.title = 'Nebula Graph 图空间级 WAL 缓存配置需求设计'
doc.core_properties.subject = '开发启动与团队需求澄清'
doc.core_properties.author = 'Codex'
doc.core_properties.keywords = 'Nebula Graph, WAL, 图空间, 内存, 需求设计'
doc.core_properties.comments = '基于固定版本源码和现有实验的评审稿，尚未实施内核改造。'
doc.save(OUTPUT)
log = {
    'generated_at': datetime.now(timezone.utc).isoformat(),
    'source': str(SOURCE), 'output': str(OUTPUT),
    'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'docx_sha256': hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables),
    'font': FONT, 'status': 'render_and_visual_review_required'
}
(ROOT/'design-evidence'/'docx-build.json').write_text(json.dumps(log, ensure_ascii=False, indent=2)+'\n')
print(json.dumps(log, ensure_ascii=False, indent=2))
